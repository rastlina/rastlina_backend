"""orders/admin.py"""

import csv

from django.contrib import admin
from django.http import HttpResponse
from django.utils import timezone
from django.utils.html import format_html
from django.contrib import messages
from .services import approve_return_request
from .services import approve_exchange_request
from .models import ExchangeCode, Order, OrderItem, OrderReturnRequest, ReturnRequest
import tempfile
import requests

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.drawing.image import Image as XLImage

def export_to_excel(modeladmin, request, queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"

    headers = [
        'Order ID',
        'Date',
        'Customer',
        'Phone',
        'Email',
        'Payment Status',
        'Order Status',
        'Payment Method',
        'Address',
        'City',
        'State',
        'Pincode',
        'SKU',
        'Product',
        'Variant',
        'Quantity',
        'Unit Price',
        'Line Total',

        'Image',
    ]

    ws.append(headers)

    # Header styling
    for cell in ws[1]:
        cell.font = Font(bold=True)

    row_num = 2

    for order in queryset.prefetch_related('items'):

        for item in order.items.all():
            sku = getattr(item, 'sku', '') or getattr(item.product, 'sku', '—')

            product_name = item.product_name

            variant = item.variant_label or '—'

            quantity = item.quantity

            unit_price = float(item.price)

            line_total = quantity * unit_price

            address = " ".join(filter(None, [
                order.shipping_address,
                order.apartment,
                order.landmark,
            ]))

            ws.cell(row=row_num, column=1, value=order.id)
            ws.cell(row=row_num, column=2, value=timezone.localtime(order.created_at).strftime('%d-%m-%Y %H:%M'))
            ws.cell(row=row_num, column=3, value=f"{order.first_name} {order.last_name}")
            ws.cell(row=row_num, column=4, value=order.phone)
            ws.cell(row=row_num, column=5, value=order.email)
            ws.cell(row=row_num, column=6, value=order.payment_status)
            ws.cell(row=row_num, column=7, value=order.order_status)
            ws.cell(row=row_num, column=8, value=order.payment_method)

            ws.cell(row=row_num, column=9, value=address)
            ws.cell(row=row_num, column=10, value=order.city)
            ws.cell(row=row_num, column=11, value=order.state)
            ws.cell(row=row_num, column=12, value=order.zip_code)

            ws.cell(row=row_num, column=13, value=sku)
            ws.cell(row=row_num, column=14, value=product_name)
            ws.cell(row=row_num, column=15, value=variant)
            ws.cell(row=row_num, column=16, value=quantity)
            ws.cell(row=row_num, column=17, value=unit_price)
            ws.cell(row=row_num, column=18, value=line_total)
            # Add image
            if item.image_url:
                try:
                    response = requests.get(item.image_url, timeout=10)

                    if response.status_code == 200:
                        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".jpg")
                        tmp.write(response.content)
                        tmp.close()

                        img = XLImage(tmp.name)
                        img.width = 60
                        img.height = 60

                        ws.add_image(img, f'S{row_num}')

                        ws.row_dimensions[row_num].height = 50

                except Exception:
                    pass

            row_num += 1

    # Column widths
    widths = {
        'A': 12,
        'B': 20,
        'C': 25,
        'D': 18,
        'E': 28,
        'F': 18,
        'G': 18,
        'H': 18,
        'I': 40,
        'J': 18,
        'K': 18,
        'L': 14,
        'M': 18,   # SKU
        'N': 35,   # Product
        'O': 25,   # Variant
        'P': 10,   # Qty
        'Q': 14,   # Unit Price
        'R': 14,   # Line Total
        'S': 18,   # Image
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = 'attachment; filename=orders.xlsx'

    wb.save(response)

    return response


export_to_excel.short_description = "Export selected orders to Excel"

class OrderItemInline(admin.TabularInline):
    model = OrderItem
    extra = 0
    readonly_fields = ['product_name', 'variant_label', 'price', 'quantity', 'image_preview']

    def image_preview(self, obj):
        if obj.image_url:
            return format_html('<img src="{}" style="width:50px;height:auto;" />', obj.image_url)
        return '—'
    image_preview.short_description = 'Image'


@admin.register(Order)
class OrderAdmin(admin.ModelAdmin):
    # Removed: email, payment_method; added phone + short_address
    list_display = [
        'id', 'full_name', 'phone', 'is_guest', 'total_amount',
        'payment_status', 'order_status', 'short_address', 'created_at',
    ]
    list_editable = ['order_status', 'payment_status']
    list_filter = ['payment_status', 'order_status', 'created_at']
    search_fields = ['id', 'first_name', 'last_name', 'email', 'phone', 'razorpay_order_id']
    readonly_fields = ['created_at', 'updated_at', 'razorpay_order_id', 'razorpay_payment_id']
    actions = [export_to_excel]
    inlines = [OrderItemInline]
    list_per_page = 25
    fieldsets = (
        ('Customer', {
            'fields': (('first_name', 'last_name'), 'email', 'phone', 'user'),
        }),
        ('Shipping', {
            'fields': (
                'shipping_address', 'apartment', 'landmark',
                'city', 'state', 'zip_code', 'country',
            ),
        }),
        ('Financials', {
            'fields': (
                'subtotal', 'discount_amount', 'shipping_fee',
                'cod_fee', 'tax_amount', 'total_amount',
                'coupon_code', 'exchange_code_used',
            ),
        }),
        ('Payment & Status', {
            'fields': (
                'payment_method', 'payment_status', 'order_status',
                'delivered_at', 'razorpay_order_id', 'razorpay_payment_id',
            ),
        }),
        ('Policy', {'fields': ('accepted_return_policy',)}),
        ('Tracking', {'fields': ('tracking_link', 'tracking_note')}),
        ('Timestamps', {'fields': ('created_at', 'updated_at')}),
    )

    def full_name(self, obj):
        return f'{obj.first_name} {obj.last_name}'.strip() or '—'
    full_name.short_description = 'Customer'

    def is_guest(self, obj):
        return obj.user is None
    is_guest.boolean = True
    is_guest.short_description = 'Guest?'

    def short_address(self, obj):
        parts = [p for p in [obj.city, obj.state] if p]
        return ', '.join(parts) or '—'
    short_address.short_description = 'Location'

    def save_model(self, request, obj, form, change):
        if change and 'order_status' in form.changed_data and obj.order_status == 'Delivered':
            if not obj.delivered_at:
                obj.delivered_at = timezone.now()
        super().save_model(request, obj, form, change)


# ─── Exchange Request admin ───────────────────────────────────────────────────

def approve_exchange_and_generate_code(modeladmin, request, queryset):
    success_count = 0

    for req in queryset:
        result = approve_exchange_request(req.id)

        if result.get('success'):
            success_count += 1
        else:
            modeladmin.message_user(
                request,
                f"Exchange #{req.id} failed: {result.get('error')}",
                level=messages.ERROR,
            )

    if success_count:
        modeladmin.message_user(
            request,
            f"{success_count} exchange request(s) approved successfully.",
            level=messages.SUCCESS,
        )


approve_exchange_and_generate_code.short_description = (
    'Approve & generate exchange codes'
)
def reject_exchange(modeladmin, request, queryset):
    queryset.filter(status='Pending').update(
        status='Rejected',
        admin_notes='Your exchange request has been reviewed and rejected. Contact support if you have questions.',
    )

reject_exchange.short_description = 'Reject selected exchange requests'


@admin.register(ReturnRequest)
class ExchangeRequestAdmin(admin.ModelAdmin):
    list_display = [
        'id', 'get_order_id', 'get_customer', 'status', 'has_video', 'created_at',
    ]
    list_filter = ['status']
    search_fields = ['order__id', 'order__email', 'order__user__email']
    readonly_fields = ['order', 'user', 'guest_email', 'created_at', 'updated_at', 'exchange_code']
    actions = [approve_exchange_and_generate_code, reject_exchange]
    list_per_page = 25

    def get_order_id(self, obj):
        return f'#{obj.order_id}'
    get_order_id.short_description = 'Order'

    def get_customer(self, obj):
        if obj.user:
            return obj.user.email
        return obj.guest_email or obj.order.email or '—'
    get_customer.short_description = 'Customer'

    def has_video(self, obj):
        return bool(obj.defect_video_url)
    has_video.boolean = True
    has_video.short_description = 'Video?'


# ─── Return Request admin ─────────────────────────────────────────────────────

def approve_return(modeladmin, request, queryset):
    """
    Approve physical return requests AND restore stock.
    """

    success_count = 0

    for req in queryset:
        result = approve_return_request(
            req.id,
            restore_stock=True,   # IMPORTANT
        )

        if result.get('success'):
            success_count += 1
        else:
            modeladmin.message_user(
                request,
                f"Return #{req.id} failed: {result.get('error')}",
                level=messages.ERROR,
            )

    if success_count:
        modeladmin.message_user(
            request,
            f"{success_count} return request(s) approved successfully.",
            level=messages.SUCCESS,
        )


approve_return.short_description = (
    'Approve selected return requests & restore stock'
)

def reject_return(modeladmin, request, queryset):
    queryset.filter(status='Pending').update(
        status='Rejected',
        admin_notes='Your return request has been reviewed and rejected. Contact support if needed.',
    )

reject_return.short_description = 'Reject selected return requests'


@admin.register(OrderReturnRequest)
class OrderReturnRequestAdmin(admin.ModelAdmin):
    list_display = [
        'id', 'get_order_id', 'get_customer', 'status',
        'refund_initiated', 'has_video', 'created_at',
    ]
    list_filter = ['status', 'refund_initiated']
    search_fields = ['order__id', 'order__email', 'order__user__email']
    readonly_fields = ['order', 'user', 'guest_email', 'created_at', 'updated_at']
    actions = [approve_return, reject_return]
    list_per_page = 25
    fieldsets = (
        ('Request', {'fields': ('order', 'user', 'guest_email', 'reason', 'video_url')}),
        ('Status', {'fields': ('status', 'admin_notes')}),
        ('Refund', {'fields': ('refund_initiated', 'razorpay_refund_id')}),
        ('Timestamps', {'fields': ('created_at', 'updated_at')}),
    )

    def get_order_id(self, obj):
        return f'#{obj.order_id}'
    get_order_id.short_description = 'Order'

    def get_customer(self, obj):
        if obj.user:
            return obj.user.email
        return obj.guest_email or obj.order.email or '—'
    get_customer.short_description = 'Customer'

    def has_video(self, obj):
        return bool(obj.video_url)
    has_video.boolean = True
    has_video.short_description = 'Video?'


# ─── Exchange Code admin ──────────────────────────────────────────────────────

@admin.register(ExchangeCode)
class ExchangeCodeAdmin(admin.ModelAdmin):
    list_display = [
        'code', 'get_order', 'original_order_value', 'is_used', 'expires_at', 'created_at',
    ]
    list_filter = ['is_used']
    search_fields = ['code', 'order__id']
    readonly_fields = ['code', 'created_at', 'used_at']
    list_per_page = 25

    def get_order(self, obj):
        return f'Order #{obj.order_id}'
    get_order.short_description = 'Order'